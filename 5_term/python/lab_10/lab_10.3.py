from docx import Document

doc = Document(
    "C:\\Users\\Бобр Даша\\Desktop\\university\\3 КУРС\\5 сем\\Язык Python для работы с данными\\lab_10\\lab_10.docx"
)

ATmega328 = {}

for table in doc.tables:
    ATmega328["Flash (1 кБ flash-памяти занят загрузчиком)"] = table.cell(1, 2).text
    ATmega328["SRAM"] = table.cell(2, 2).text
    ATmega328["EEPROM"] = table.cell(3, 2).text

print(ATmega328)
